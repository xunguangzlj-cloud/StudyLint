from pathlib import Path

import pymupdf
from docx import Document

from studylint.fulltext import ReferenceResolution
from studylint.manuscripts import (
    audit_manuscript,
    discover_reference_files,
    extract_citation_claims,
    scan_manuscript_issues,
)


def make_pdf(path: Path, text: str) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(path)
    document.close()


def test_extracts_numeric_citation_groups_and_ranges() -> None:
    blocks = [
        "The first claim [1, 2]. Another claim [3-4].",
        "# References",
        "[1] First paper.",
        "[2] Second paper.",
        "[3] Third paper.",
        "[4] Fourth paper.",
    ]

    claims, references = extract_citation_claims(blocks)

    assert [(number, claim) for _, number, claim in claims] == [
        (1, "The first claim."),
        (2, "The first claim."),
        (3, "Another claim."),
        (4, "Another claim."),
    ]
    assert references[3] == "Third paper."


def test_extracts_fullwidth_parenthesized_and_superscript_citations() -> None:
    blocks = [
        "全角方括号引用［1，2］。圆括号引用（3）。上标引用⁴。",
        "参考文献",
        "［1］第一篇论文。",
        "［2］第二篇论文。",
        "［3］第三篇论文。",
        "［4］第四篇论文。",
    ]

    claims, references = extract_citation_claims(blocks)

    assert [(number, claim) for _, number, claim in claims] == [
        (1, "全角方括号引用。"),
        (2, "全角方括号引用。"),
        (3, "圆括号引用。"),
        (4, "上标引用。"),
    ]
    assert references[4] == "第四篇论文。"


def test_extracts_author_year_citations_with_unnumbered_references() -> None:
    blocks = [
        "已有研究发现该方法有效（张三，2024；Smith et al., 2023）。",
        "李四（2022）进一步验证了该结论。",
        "参考文献",
        "张三. 该方法的实证研究. 2024.",
        "Smith, J., Doe, A. Evidence for the method. 2023.",
        "李四. 后续验证研究. 2022.",
    ]

    claims, references = extract_citation_claims(blocks)

    assert [(number, claim) for _, number, claim in claims] == [
        (1, "已有研究发现该方法有效。"),
        (2, "已有研究发现该方法有效。"),
        (3, "进一步验证了该结论。"),
    ]
    assert references[1].startswith("张三")
    assert references[2].startswith("Smith")
    assert references[3].startswith("李四")


def test_author_year_citation_is_not_guessed_when_reference_is_ambiguous() -> None:
    blocks = [
        "该方法有效（Smith, 2023）。",
        "References",
        "Smith, J. First study. 2023.",
        "Smith, A. Second study. 2023.",
    ]

    claims, _ = extract_citation_claims(blocks)

    assert claims == []


def test_docx_superscript_citation_is_recognized(tmp_path: Path) -> None:
    manuscript = tmp_path / "draft.docx"
    document = Document()
    paragraph = document.add_paragraph("该结论得到研究支持")
    citation = paragraph.add_run("1")
    citation.font.superscript = True
    document.add_paragraph("参考文献")
    document.add_paragraph("[1] Example study. 2024.")
    document.save(manuscript)

    audit = audit_manuscript(manuscript, [tmp_path])[0]

    assert audit.citation_number == 1
    assert audit.claim == "该结论得到研究支持"
    assert audit.location == "第1段"


def test_audits_claim_against_numbered_reference_pdf(tmp_path: Path) -> None:
    manuscript = tmp_path / "draft.md"
    source = tmp_path / "1-treatment-study.pdf"
    claim = "The treatment reduced complications by 20 percent."
    manuscript.write_text(
        f"{claim} [1]\n\n# References\n[1] Example. Treatment study. 2024.\n",
        encoding="utf-8",
    )
    make_pdf(source, claim)

    audit = audit_manuscript(manuscript, [source])[0]

    assert audit.verdict == "DIRECT_SUPPORT"
    assert audit.citation_number == 1
    assert audit.source_path == str(source.resolve())
    assert audit.evidences[0].page == 1


def test_markdown_audit_reports_exact_source_line(tmp_path: Path) -> None:
    manuscript = tmp_path / "draft.md"
    manuscript.write_text(
        "# 标题\n\n需要核验的结论[1]。\n\n参考文献\n[1] Example study. 2024.\n",
        encoding="utf-8",
    )

    audit = audit_manuscript(manuscript, [tmp_path])[0]

    assert audit.location == "第3行"


def test_missing_reference_and_unavailable_source_are_separate(tmp_path: Path) -> None:
    manuscript = tmp_path / "draft.txt"
    source_dir = tmp_path / "sources"
    source_dir.mkdir()
    manuscript.write_text(
        "A supported-looking claim [1]. A dangling claim [2].\n"
        "References\n"
        "[1] Existing reference without a supplied PDF.\n",
        encoding="utf-8",
    )

    audits = audit_manuscript(manuscript, [source_dir])

    assert audits[0].verdict == "SOURCE_UNAVAILABLE"
    assert audits[1].verdict == "MISSING_REFERENCE"


def test_repeated_claim_is_audited_for_each_citation(tmp_path: Path) -> None:
    manuscript = tmp_path / "draft.md"
    first_source = tmp_path / "1-first-study.pdf"
    second_source = tmp_path / "2-second-study.pdf"
    claim = "The treatment reduced complications by 20 percent."
    manuscript.write_text(
        f"{claim} [1]\n\n{claim} [2]\n\n"
        "# References\n[1] First study.\n[2] Second study.\n",
        encoding="utf-8",
    )
    make_pdf(first_source, claim)
    make_pdf(second_source, claim)

    audits = audit_manuscript(manuscript, [first_source, second_source])

    assert len(audits) == 2
    assert [audit.citation_number for audit in audits] == [1, 2]
    assert all(audit.verdict == "DIRECT_SUPPORT" for audit in audits)


def test_reference_discovery_accepts_pdf_and_epub(tmp_path: Path) -> None:
    pdf = tmp_path / "1-paper.pdf"
    epub = tmp_path / "2-book.epub"
    ignored = tmp_path / "notes.txt"
    pdf.write_bytes(b"%PDF-test")
    epub.write_bytes(b"epub-test")
    ignored.write_text("ignored", encoding="utf-8")

    found = discover_reference_files([tmp_path])

    assert found == [pdf.resolve(), epub.resolve()]


def test_scans_high_confidence_document_level_issues(tmp_path: Path) -> None:
    manuscript = tmp_path / "draft.md"
    manuscript.write_text(
        "The sample included 100 students [1].\n"
        "The sample included 120 students [1].\n\n"
        "# References\n"
        "[1] Real paper. 2024.\n"
        "[2] TODO citation. 2099.\n",
        encoding="utf-8",
    )

    issues = scan_manuscript_issues(manuscript)
    codes = {issue.code for issue in issues}

    assert "UNUSED_REFERENCE" in codes
    assert "PLACEHOLDER_REFERENCE" in codes
    assert "FUTURE_REFERENCE_YEAR" in codes
    assert "INCONSISTENT_NUMERIC_STATEMENT" in codes


def test_citation_numbers_and_years_are_not_treated_as_result_data(tmp_path: Path) -> None:
    manuscript = tmp_path / "draft.md"
    manuscript.write_text(
        "该方法能够改善学习效果［1］。\n"
        "该方法能够改善学习效果（张三，2024）。\n\n"
        "参考文献\n"
        "［1］李四. 第一项研究. 2023.\n"
        "［2］张三. 第二项研究. 2024.\n",
        encoding="utf-8",
    )

    issues = scan_manuscript_issues(manuscript)

    assert all(issue.code != "INCONSISTENT_NUMERIC_STATEMENT" for issue in issues)


def test_auto_fetch_uses_resolved_open_fulltext(
    tmp_path: Path, monkeypatch
) -> None:
    manuscript = tmp_path / "draft.md"
    downloaded = tmp_path / "cache" / "reference-1.pdf"
    downloaded.parent.mkdir()
    claim = "The treatment reduced complications by 20 percent."
    manuscript.write_text(
        f"{claim} [1]\n\n# References\n[1] Example study. 2024.\n",
        encoding="utf-8",
    )
    make_pdf(downloaded, claim)

    monkeypatch.setattr(
        "studylint.manuscripts.resolve_reference_sources",
        lambda references, local_matches, cache_dir: {
            1: ReferenceResolution(
                number=1,
                reference=references[1],
                verification_verdict="VERIFIED_METADATA",
                verification_message="题名与论文元数据高置信匹配。",
                url="https://doi.org/10.1234/example",
                work_type="journal-article",
                is_retracted=False,
                full_text_url="https://example.org/paper.pdf",
                license="cc-by",
                path=downloaded,
                source_status="DOWNLOADED",
                source_message="已获取开放原文。",
            )
        },
    )

    audit = audit_manuscript(
        manuscript, [], auto_fetch=True, cache_dir=tmp_path / "cache"
    )[0]

    assert audit.verdict == "DIRECT_SUPPORT"
    assert audit.reference_verdict == "VERIFIED_METADATA"
    assert audit.source_status == "DOWNLOADED"
    assert audit.source_path == str(downloaded)


def test_auto_fetch_unavailable_requests_local_pdf_or_epub(
    tmp_path: Path, monkeypatch
) -> None:
    manuscript = tmp_path / "draft.md"
    manuscript.write_text(
        "A claim [1].\n\n# References\n[1] Example study. 2024.\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(
        "studylint.manuscripts.resolve_reference_sources",
        lambda references, local_matches, cache_dir: {
            1: ReferenceResolution(
                number=1,
                reference=references[1],
                verification_verdict="VERIFIED_METADATA",
                verification_message="题名与论文元数据高置信匹配。",
                url="https://doi.org/10.1234/example",
                work_type="journal-article",
                is_retracted=False,
                full_text_url="",
                license="",
                path=None,
                source_status="UNAVAILABLE",
                source_message="当前未取得可核验开放全文；这不代表论文不存在。",
            )
        },
    )

    audit = audit_manuscript(manuscript, [], auto_fetch=True)[0]

    assert audit.verdict == "SOURCE_UNAVAILABLE"
    assert "PDF或EPUB" in audit.explanation
    assert "不代表论文不存在" in audit.explanation
