from pathlib import Path
from zipfile import ZipFile

from studylint.parsers import (
    discover_sources,
    extract_citations,
    load_source,
    load_sources,
    parse_notes,
    parse_time,
)


def test_parse_time() -> None:
    assert parse_time("00:01:30") == 90
    assert parse_time("01:02:03,400") == 3723
    assert parse_time("00:70:00") is None


def test_extract_citations() -> None:
    citations = extract_citations(
        "内容 [slides.pdf#page=12] [lecture.srt#time=00:31:42]"
    )
    assert [(item.source_name, item.locator_type, item.locator_value) for item in citations] == [
        ("slides.pdf", "page", 12),
        ("lecture.srt", "time", 1902),
    ]

    aliases = extract_citations(
        "内容 [slides.pdf#页码=3] [lecture.srt#t=00:00:10]"
    )
    assert [(item.locator_type, item.locator_value) for item in aliases] == [
        ("page", 3),
        ("time", 10),
    ]


def test_markdown_page_markers(tmp_path: Path) -> None:
    source = tmp_path / "slides.md"
    source.write_text("第一页\n<!-- page: 3 -->\n第三页", encoding="utf-8")
    document = load_source(source)
    assert [(span.locator_value, span.text) for span in document.spans] == [
        (1, "第一页"),
        (3, "第三页"),
    ]


def test_srt_segments(tmp_path: Path) -> None:
    source = tmp_path / "lecture.srt"
    source.write_text(
        "1\n00:00:01,000 --> 00:00:04,000\n第一句话\n\n"
        "2\n00:00:05,000 --> 00:00:08,000\n第二句话\n",
        encoding="utf-8",
    )
    document = load_source(source)
    assert len(document.spans) == 2
    assert document.spans[1].locator_value == 5
    assert document.spans[1].end_value == 8


def test_pdf_pages(tmp_path: Path) -> None:
    import pymupdf

    source = tmp_path / "slides.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Page one evidence")
    document.save(source)
    document.close()

    parsed = load_source(source)
    assert len(parsed.spans) == 1
    assert parsed.spans[0].locator_value == 1
    assert "Page one evidence" in parsed.spans[0].text


def test_pptx_slides(tmp_path: Path) -> None:
    from pptx import Presentation

    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Slide evidence"
    presentation.save(source)

    parsed = load_source(source)
    assert len(parsed.spans) == 1
    assert parsed.spans[0].locator_value == 1
    assert "Slide evidence" in parsed.spans[0].text


def test_docx_notes(tmp_path: Path) -> None:
    from docx import Document

    notes = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("信息能够减少不确定性。[slides.pdf#page=2]")
    document.save(notes)

    units = parse_notes(notes)
    assert len(units) == 1
    assert units[0].citations[0].locator_value == 2


def test_pdf_notes(tmp_path: Path) -> None:
    import pymupdf

    notes = tmp_path / "notes.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "A summary statement")
    document.save(notes)
    document.close()

    units = parse_notes(notes)

    assert [unit.text for unit in units] == ["A summary statement"]


def test_epub_source_and_notes(tmp_path: Path) -> None:
    book = tmp_path / "book.epub"
    with ZipFile(book, "w") as archive:
        archive.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>',
        )
        archive.writestr(
            "OEBPS/content.opf",
            '<package xmlns="http://www.idpf.org/2007/opf"><manifest>'
            '<item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/>'
            '</manifest><spine><itemref idref="chapter"/></spine></package>',
        )
        archive.writestr(
            "OEBPS/chapter.xhtml",
            "<html><body><h1>Chapter one</h1><p>Evidence sentence.</p></body></html>",
        )

    parsed = load_source(book)
    units = parse_notes(book)

    assert parsed.spans[0].locator_type == "chapter"
    assert "Evidence sentence." in parsed.spans[0].text
    assert any("Evidence sentence." in unit.text for unit in units)


def test_discover_sources_excludes_notes(tmp_path: Path) -> None:
    notes = tmp_path / "notes.md"
    notes.write_text("笔记", encoding="utf-8")
    (tmp_path / "slides.pdf").write_bytes(b"not parsed in this test")
    nested = tmp_path / "nested"
    nested.mkdir()
    (nested / "lecture.srt").write_text("", encoding="utf-8")
    (nested / "ignored.csv").write_text("", encoding="utf-8")

    found = discover_sources(tmp_path, exclude=notes)
    assert [path.name for path in found] == ["lecture.srt", "slides.pdf"]


def test_load_sources_preserves_input_order(tmp_path: Path) -> None:
    first = tmp_path / "first.md"
    second = tmp_path / "second.md"
    first.write_text("第一份", encoding="utf-8")
    second.write_text("第二份", encoding="utf-8")

    documents = load_sources([second, first])
    assert [document.name for document in documents] == ["second.md", "first.md"]
