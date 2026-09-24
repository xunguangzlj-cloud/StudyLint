from __future__ import annotations

import posixpath
import re
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote
from zipfile import BadZipFile, ZipFile

from studylint.models import Citation, NoteUnit, SourceDocument, SourceSpan


CITATION_PATTERN = re.compile(
    r"\[(?P<source>[^\]#]+)#(?P<kind>page|p|页|页码|time|t|时间)=(?P<value>[^\]]+)\]",
    re.I,
)
PAGE_MARKER_PATTERN = re.compile(r"^\s*<!--\s*page\s*:\s*(\d+)\s*-->\s*$", re.I)
SRT_TIME_PATTERN = re.compile(
    r"(?P<start>\d{2}:\d{2}:\d{2}[,.]\d{3})\s*-->\s*"
    r"(?P<end>\d{2}:\d{2}:\d{2}[,.]\d{3})"
)
IGNORE_PATTERN = re.compile(
    r"^\s*<!--\s*studylint-ignore\s+([A-Za-z0-9_, -]+)\s*-->\s*$",
    re.I,
)
SUPPORTED_SOURCE_SUFFIXES = {
    ".md", ".txt", ".pdf", ".epub", ".pptx", ".srt", ".docx"
}
SUPPORTED_NOTE_SUFFIXES = {".md", ".txt", ".docx", ".pdf", ".epub"}


def parse_time(value: str) -> int | None:
    match = re.fullmatch(r"\s*(\d{1,2}):(\d{2}):(\d{2})(?:[,.](\d{3}))?\s*", value)
    if not match:
        return None
    hours, minutes, seconds = (int(part) for part in match.group(1, 2, 3))
    if minutes >= 60 or seconds >= 60:
        return None
    return hours * 3600 + minutes * 60 + seconds


def extract_citations(text: str) -> tuple[Citation, ...]:
    citations: list[Citation] = []
    for match in CITATION_PATTERN.finditer(text):
        raw_kind = match.group("kind").lower()
        kind = "page" if raw_kind in {"page", "p", "页", "页码"} else "time"
        raw_value = match.group("value").strip()
        if kind == "page":
            try:
                value = int(raw_value)
            except ValueError:
                value = None
        else:
            value = parse_time(raw_value)
        citations.append(
            Citation(
                source_name=match.group("source").strip(),
                locator_type=kind,
                locator_value=value,
                raw_value=raw_value,
                raw=match.group(0),
            )
        )
    return tuple(citations)


def _parse_text_notes(path: Path) -> list[NoteUnit]:
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    units: list[NoteUnit] = []
    in_code_block = False
    ignored_codes: set[str] = set()
    for line_number, line in enumerate(lines, start=1):
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code_block = not in_code_block
            continue
        ignore_match = IGNORE_PATTERN.match(stripped)
        if ignore_match:
            ignored_codes.update(
                code.strip().upper()
                for code in re.split(r"[,\s]+", ignore_match.group(1))
                if code.strip()
            )
            continue
        if (
            in_code_block
            or not stripped
            or stripped.startswith("#")
            or stripped.startswith("<!--")
        ):
            continue
        units.append(
            NoteUnit(
                line=line_number,
                text=stripped,
                citations=extract_citations(stripped),
                ignored_codes=tuple(sorted(ignored_codes)),
            )
        )
        ignored_codes.clear()
    return units


def _parse_docx_notes(path: Path) -> list[NoteUnit]:
    from docx import Document

    document = Document(path)
    units: list[NoteUnit] = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        units.append(
            NoteUnit(
                line=paragraph_number,
                text=text,
                citations=extract_citations(text),
            )
        )
    return units


def _parse_document_notes(path: Path) -> list[NoteUnit]:
    document = load_source(path)
    units: list[NoteUnit] = []
    unit_number = 0
    for span in document.spans:
        for raw_text in span.text.splitlines():
            text = raw_text.strip()
            if not text:
                continue
            unit_number += 1
            units.append(
                NoteUnit(
                    line=unit_number,
                    text=text,
                    citations=extract_citations(text),
                )
            )
    return units


def parse_notes(path: Path) -> list[NoteUnit]:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return _parse_text_notes(path)
    if suffix == ".docx":
        return _parse_docx_notes(path)
    if suffix in {".pdf", ".epub"}:
        return _parse_document_notes(path)
    supported = ", ".join(sorted(SUPPORTED_NOTE_SUFFIXES))
    raise ValueError(f"不支持的AI总结格式：{suffix or '(无扩展名)'}。支持：{supported}")


def _parse_markdown_or_text(path: Path) -> SourceDocument:
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    page = 1
    pages: dict[int, list[str]] = {page: []}
    for line in lines:
        marker = PAGE_MARKER_PATTERN.match(line)
        if marker:
            page = int(marker.group(1))
            pages.setdefault(page, [])
            continue
        pages[page].append(line)

    spans = [
        SourceSpan(path.name, "page", number, "\n".join(content).strip())
        for number, content in sorted(pages.items())
    ]
    return SourceDocument(path.name, path, path.suffix.lower().lstrip("."), spans)


def _parse_pdf(path: Path) -> SourceDocument:
    import pymupdf

    document = pymupdf.open(path)
    try:
        spans = [
            SourceSpan(path.name, "page", index + 1, page.get_text("text").strip())
            for index, page in enumerate(document)
        ]
    finally:
        document.close()
    return SourceDocument(path.name, path, "pdf", spans)


class _EpubTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.ignored_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag.lower() in {"script", "style", "nav"}:
            self.ignored_depth += 1
        elif tag.lower() in {"p", "div", "h1", "h2", "h3", "li", "br"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "nav"} and self.ignored_depth:
            self.ignored_depth -= 1
        elif tag.lower() in {"p", "div", "h1", "h2", "h3", "li"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.ignored_depth and data.strip():
            self.parts.append(data)

    def text(self) -> str:
        lines = (
            re.sub(r"\s+", " ", line).strip()
            for line in "".join(self.parts).splitlines()
        )
        return "\n".join(line for line in lines if line)


def _parse_epub(path: Path) -> SourceDocument:
    try:
        with ZipFile(path) as archive:
            container = ET.fromstring(archive.read("META-INF/container.xml"))
            rootfile = container.find(".//{*}rootfile")
            if rootfile is None or not rootfile.get("full-path"):
                raise ValueError("EPUB缺少内容清单。")
            package_path = rootfile.get("full-path", "")
            package = ET.fromstring(archive.read(package_path))
            package_dir = posixpath.dirname(package_path)
            manifest = {
                item.get("id", ""): posixpath.normpath(
                    posixpath.join(
                        package_dir,
                        unquote(item.get("href", "").split("#", 1)[0]),
                    )
                )
                for item in package.findall(".//{*}manifest/{*}item")
                if item.get("id") and item.get("href")
            }
            chapter_paths = [
                manifest.get(item.get("idref", ""), "")
                for item in package.findall(".//{*}spine/{*}itemref")
            ]
            spans: list[SourceSpan] = []
            for chapter_path in chapter_paths:
                if not chapter_path:
                    continue
                parser = _EpubTextParser()
                parser.feed(
                    archive.read(chapter_path).decode("utf-8", errors="replace")
                )
                text = parser.text()
                if text:
                    spans.append(
                        SourceSpan(path.name, "chapter", len(spans) + 1, text)
                    )
    except (BadZipFile, KeyError, ET.ParseError) as error:
        raise ValueError("EPUB文件损坏或结构不受支持。") from error
    if not spans:
        raise ValueError("EPUB中没有可提取的正文。")
    return SourceDocument(path.name, path, "epub", spans)


def _parse_pptx(path: Path) -> SourceDocument:
    from pptx import Presentation

    presentation = Presentation(path)
    spans: list[SourceSpan] = []
    for index, slide in enumerate(presentation.slides, start=1):
        text_parts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        spans.append(SourceSpan(path.name, "page", index, "\n".join(text_parts)))
    return SourceDocument(path.name, path, "pptx", spans)


def _parse_docx_source(path: Path) -> SourceDocument:
    from docx import Document

    document = Document(path)
    text = "\n".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )
    return SourceDocument(
        path.name,
        path,
        "docx",
        [SourceSpan(path.name, "page", 1, text)],
    )


def _parse_srt(path: Path) -> SourceDocument:
    content = path.read_text(encoding="utf-8-sig")
    blocks = re.split(r"\r?\n\s*\r?\n", content.strip())
    spans: list[SourceSpan] = []
    for block in blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        time_index = next(
            (index for index, line in enumerate(lines) if SRT_TIME_PATTERN.fullmatch(line)),
            None,
        )
        if time_index is None:
            continue
        match = SRT_TIME_PATTERN.fullmatch(lines[time_index])
        if match is None:
            continue
        start = parse_time(match.group("start"))
        end = parse_time(match.group("end"))
        if start is None or end is None:
            continue
        text = " ".join(lines[time_index + 1 :]).strip()
        spans.append(SourceSpan(path.name, "time", start, text, end))
    return SourceDocument(path.name, path, "srt", spans)


def load_source(path: Path) -> SourceDocument:
    suffix = path.suffix.lower()
    parsers = {
        ".md": _parse_markdown_or_text,
        ".txt": _parse_markdown_or_text,
        ".pdf": _parse_pdf,
        ".epub": _parse_epub,
        ".pptx": _parse_pptx,
        ".srt": _parse_srt,
        ".docx": _parse_docx_source,
    }
    if suffix not in parsers:
        supported = ", ".join(sorted(parsers))
        raise ValueError(f"Unsupported source type: {suffix or '(none)'}. Supported: {supported}")
    return parsers[suffix](path)


def load_sources(paths: list[Path]) -> list[SourceDocument]:
    if len(paths) <= 1:
        return [load_source(path) for path in paths]
    workers = min(8, len(paths))
    with ThreadPoolExecutor(max_workers=workers) as executor:
        return list(executor.map(load_source, paths))


def discover_sources(directory: Path, exclude: Path | None = None) -> list[Path]:
    if not directory.is_dir():
        raise ValueError(f"Source directory not found: {directory}")
    excluded = exclude.resolve() if exclude else None
    return sorted(
        path
        for path in directory.rglob("*")
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_SOURCE_SUFFIXES
        and (excluded is None or path.resolve() != excluded)
    )
